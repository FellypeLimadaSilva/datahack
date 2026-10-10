from __future__ import annotations

import argparse
import csv
import gzip
import io
import json
import random
import sqlite3
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook

FIRST = ["Ana", "Bruno", "Carla", "Diego", "Elisa", "Fábio", "Gabriela", "Hélio", "Íris", "João"]
LAST = ["Silva", "Souza", "Conceição", "Araújo", "Gonçalves", "Magalhães", "Brandão", "Peçanha"]
UFS = ["MT", "MS", "GO", "SP", "PR", "BA"]


def cpf(rng: random.Random) -> str:
    base = [rng.randint(0, 9) for _ in range(9)]
    for size in (9, 10):
        total = sum(d * w for d, w in zip(base, range(size + 1, 1, -1), strict=True))
        base.append((total * 10) % 11 % 10)
    s = "".join(map(str, base))
    return f"{s[:3]}.{s[3:6]}.{s[6:9]}-{s[9:]}"


def brl(value: float) -> str:
    whole, cents = f"{value:,.2f}".split(".")
    return f"R$ {whole.replace(',', '.')},{cents}"


def clientes(out: Path, rng: random.Random) -> None:
    folder = out / "clientes"
    folder.mkdir(parents=True, exist_ok=True)
    header = [
        "Código",
        "Nome Completo",
        "CPF",
        "E-mail",
        "Data Nascimento",
        "Renda Mensal",
        "Ativo",
        "CEP",
        "Telefone",
    ]

    def row(i: int, bonus: float = 0) -> list[str]:
        nome = f"{rng.choice(FIRST)} {rng.choice(LAST)}"
        nasc = date(1960, 1, 1) + timedelta(days=rng.randint(0, 20000))
        return [
            f"{i:03d}",
            nome,
            cpf(rng),
            f"cliente{i}@exemplo.com.br",
            nasc.strftime("%d/%m/%Y"),
            brl(rng.uniform(1200, 25000) + bonus),
            rng.choice(["Sim", "Não"]),
            f"78{rng.randint(0, 999):03d}-{rng.randint(0, 999):03d}",
            f"(65) 9{rng.randint(1000, 9999)}-{rng.randint(1000, 9999)}",
        ]

    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(header)
    for i in range(1, 51):
        w.writerow(row(i))
    (folder / "clientes_2026_01.csv").write_bytes(buf.getvalue().encode("cp1252"))

    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\n")
    w.writerow([*header, "Segmento"])
    for i in range(40, 61):
        w.writerow([*row(i, bonus=1000), rng.choice(["Varejo", "Atacado"])])
    (folder / "clientes_2026_02.csv.gz").write_bytes(gzip.compress(buf.getvalue().encode("utf-8")))


def pedidos(out: Path, rng: random.Random, n: int) -> None:
    lines = []
    start = datetime(2026, 1, 1, 8)
    for i in range(1, n + 1):
        itens = [
            {"sku": f"SKU-{rng.randint(1, 50):04d}", "qtd": rng.randint(1, 5)}
            for _ in range(rng.randint(1, 3))
        ]
        lines.append(
            json.dumps(
                {
                    "id": i,
                    "cliente": {"codigo": f"{rng.randint(1, 60):03d}", "uf": rng.choice(UFS)},
                    "itens": itens,
                    "total": round(rng.uniform(5, 900), 2),
                    "criado_em": (start + timedelta(minutes=17 * i)).isoformat() + "-04:00",
                    "status": rng.choice(["pago", "pendente", "cancelado"]),
                },
                ensure_ascii=False,
            )
        )
    (out / "pedidos.jsonl.gz").write_bytes(gzip.compress("\n".join(lines).encode("utf-8")))


def estoque(out: Path, rng: random.Random) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(["SKU", "Descrição", "Preço", "Ativo", "Atualizado em"])
    for i in range(1, 51):
        ws.append(
            [
                f"SKU-{i:04d}",
                f"Produto {i}",
                round(rng.uniform(1, 500), 2),
                rng.random() > 0.1,
                datetime(2026, 2, 1, 9) + timedelta(hours=i),
            ]
        )
    mv = wb.create_sheet("Movimentos")
    mv.append(["Data", "SKU", "Quantidade", "Tipo"])
    for i in range(200):
        mv.append(
            [
                date(2026, 1, 1) + timedelta(days=i % 90),
                f"SKU-{rng.randint(1, 50):04d}",
                rng.randint(1, 40),
                rng.choice(["entrada", "saida"]),
            ]
        )
    wb.save(out / "Estoque.xlsx")


def notas(out: Path, rng: random.Random) -> None:
    parts = ['<?xml version="1.0" encoding="UTF-8"?>']
    parts.append('<nfeLote xmlns="http://www.portalfiscal.inf.br/nfe" versao="4.00">')
    for i in range(1, 31):
        chave = "".join(str(rng.randint(0, 9)) for _ in range(44))
        parts.append(
            f'<nota Id="NFe{i}"><chave>{chave}</chave>'
            f"<emitente><cnpj>{rng.randint(10**13, 10**14 - 1)}</cnpj>"
            f"<nome>Fornecedor {i} &amp; Cia</nome></emitente>"
            f"<valor>{rng.uniform(100, 9000):.2f}</valor>"
            f"<emissao>2026-03-{(i % 28) + 1:02d}T10:00:00</emissao></nota>"
        )
    parts.append("</nfeLote>")
    (out / "notas_fiscais.xml").write_text("".join(parts), encoding="utf-8")


def eventos(out: Path) -> None:
    folder = out / "eventos"
    folder.mkdir(parents=True, exist_ok=True)
    rows = [
        ["2026-04-01 10:00:00", "login", "usuario entrou, via app"],
        ["2026-04-01 10:05:00", "erro", 'mensagem com "aspas" e\nquebra de linha'],
        ["2026-04-01 10:05:00", "erro", 'mensagem com "aspas" e\nquebra de linha'],
        ["2026-04-01 11:00:00", "logout", "fim"],
    ]
    for name, chunk in (("log_1.csv", rows[:3]), ("log_2.csv", rows[1:])):
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\n")
        w.writerow(["quando", "tipo", "detalhe"])
        w.writerows(chunk)
        (folder / name).write_bytes(b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8"))


def export_zip(out: Path) -> None:
    content = (
        "Título\tVencimento\tValor\n"
        "Duplicata Açougue\t10/05/2026\t1.500,00\n"
        "Boleto Ótica\t5/6/2026\t89,90\n"
    )
    with zipfile.ZipFile(out / "export_erp.zip", "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("relatorios/titulos.txt", content.encode("latin-1"))
        zf.writestr("relatorios/leia.pdf", b"%PDF-1.4")


def catalogo(out: Path) -> None:
    payload = {
        "meta": {"gerado_em": "2026-05-01", "versao": 3},
        "data": {
            "items": [
                {"id": i, "nome": f"Item {i}", "atributos": {"cor": "azul", "peso_kg": i / 10}}
                for i in range(1, 21)
            ]
        },
    }
    (out / "catalogo.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def legado(out: Path) -> None:
    path = out / "legado.db"
    path.unlink(missing_ok=True)
    with sqlite3.connect(path) as conn:
        conn.execute("CREATE TABLE fornecedores (id INTEGER PRIMARY KEY, razao TEXT, uf TEXT)")
        conn.execute(
            'CREATE TABLE "Contas a Pagar" (id INTEGER, fornecedor_id INTEGER, valor REAL)'
        )
        conn.executemany(
            "INSERT INTO fornecedores VALUES (?, ?, ?)",
            [(i, f"Fornecedor {i}", UFS[i % len(UFS)]) for i in range(1, 11)],
        )
        conn.executemany(
            'INSERT INTO "Contas a Pagar" VALUES (?, ?, ?)',
            [(i, (i % 10) + 1, i * 10.5) for i in range(1, 31)],
        )


def reservadas(out: Path) -> None:
    content = "order;select;user;valor;valor;\n1;a;x;10;20;z\n2;b;y;11;21;w\n"
    (out / "reservadas.csv").write_text(content, encoding="utf-8")


def transacoes(out: Path, rng: random.Random, n: int) -> None:
    folder = out / "transacoes"
    folder.mkdir(parents=True, exist_ok=True)
    base = datetime(2026, 1, 1)
    table = pa.table(
        {
            "transacao_id": pa.array(range(1, n + 1), pa.int64()),
            "valor": pa.array([round(rng.uniform(1, 1000), 2) for _ in range(n)], pa.float64()),
            "momento": pa.array([base + timedelta(seconds=37 * i) for i in range(n)]),
            "cartao": pa.array([f"4111 1111 1111 {rng.randint(0, 9999):04d}" for _ in range(n)]),
        }
    )
    pq.write_table(table, folder / "parte_001.parquet")


CURSOS = [
    ("5101", "Direito", 1),
    ("5102", "Enfermagem", 2),
    ("5103", "Pedagogia", 1),
    ("5104", "Engenharia Civil", 2),
    ("5105", "Letras", 1),
    ("5106", "Medicina Veterinária", 2),
]


def trajetoria(out: Path, rng: random.Random) -> None:
    folder = out / "trajetoria"
    folder.mkdir(parents=True, exist_ok=True)
    header = [
        "CO_IES",
        "NO_IES",
        "TP_CATEGORIA_ADMINISTRATIVA",
        "CO_CURSO",
        "NO_CURSO",
        "CO_MUNICIPIO",
        "TP_MODALIDADE_ENSINO",
        "NU_ANO_INGRESSO",
        "NU_ANO_REFERENCIA",
        "QT_INGRESSANTE",
        "QT_DESISTENCIA",
        "QT_CONCLUINTE",
        "TAP",
        "TCA",
        "TDA",
    ]
    for coorte in (2015, 2016):
        wb = Workbook()
        ws = wb.active
        ws.title = "Indicadores"
        for line in (
            "Ministério da Educação",
            "Instituto Nacional de Estudos e Pesquisas Educacionais Anísio Teixeira",
            None,
            f"Indicadores de Trajetória da Educação Superior - Coorte {coorte}",
            "Mato Grosso",
            None,
            None,
            None,
        ):
            ws.append([line])
        ws.append(header)
        for co_curso, nome, rede in CURSOS:
            ingressantes = rng.choice([6, 8, 45, 80, 120]) if co_curso != "5101" else 150
            desist = 0
            concl = 0
            for ano in range(coorte, 2025):
                d = rng.randint(0, max(1, (ingressantes - desist - concl) // 5))
                c = rng.randint(0, 6) if ano - coorte >= 4 else 0
                desist += d
                concl += c
                ws.append(
                    [
                        f"{100 + rede}",
                        f"Instituição {rede}",
                        rede,
                        co_curso,
                        nome,
                        "5103403",
                        1,
                        coorte,
                        ano,
                        ingressantes,
                        desist,
                        concl,
                        round(100 * (ingressantes - desist - concl) / ingressantes, 2),
                        round(100 * concl / ingressantes, 2),
                        round(100 * desist / ingressantes, 2),
                    ]
                )
        ws.append([None])
        ws.append(["Fonte: Inep, Censo da Educação Superior."])
        ws.append(["Notas: TAP, TCA e TDA em percentual."])
        wb.save(folder / f"indicadores_trajetoria_educacao_superior_{coorte}_2024.xlsx")


def censo(out: Path, rng: random.Random) -> None:
    for ano in (2023, 2024):
        cursos = [
            "NU_ANO_CENSO;SG_UF;CO_IES;CO_CURSO;NO_CURSO;TP_REDE;TP_MODALIDADE_ENSINO;QT_MAT;QT_ING;QT_SIT_DESVINCULADO"
        ]
        for co_curso, nome, rede in CURSOS:
            cursos.append(
                f"{ano};MT;{100 + rede};{co_curso};{nome};{rede};1;"
                f"{rng.randint(50, 900)};{rng.randint(10, 200)};{rng.randint(0, 60)}"
            )
        ies = ["NU_ANO_CENSO;CO_IES;NO_IES;CO_MUNICIPIO;TP_CATEGORIA_ADMINISTRATIVA"]
        ies += [
            f"{ano};101;Instituição Pública;5103403;1",
            f"{ano};102;Instituição Privada;5108402;4",
        ]
        base = f"microdados_censo_da_educacao_superior_{ano}"
        with zipfile.ZipFile(out / f"{base}.zip", "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(
                f"{base}/dados/MICRODADOS_CADASTRO_CURSOS_{ano}.CSV",
                "\n".join(cursos).encode("latin-1"),
            )
            zf.writestr(
                f"{base}/dados/MICRODADOS_ED_SUP_IES_{ano}.CSV", "\n".join(ies).encode("latin-1")
            )
            zf.writestr(f"{base}/Anexos/ANEXO I/dicionário_dados_educação_superior.xlsx", b"PK")
            zf.writestr(f"{base}/leia-me/Leia-me.pdf", b"%PDF-1.4")


def noise(out: Path) -> None:
    (out / "LEIA-ME.pdf").write_bytes(b"%PDF-1.4")
    (out / "~$Estoque.xlsx").write_bytes(b"lock")
    (out / ".DS_Store").write_bytes(b"\x00")


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Gera arquivos bagunçados na inbox para demonstração.")
    p.add_argument("--out", default="data/landing/inbox")
    p.add_argument("--rows", type=int, default=50_000)
    p.add_argument("--seed", type=int, default=42)
    args = p.parse_args(argv)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    rng = random.Random(args.seed)
    clientes(out, rng)
    pedidos(out, rng, max(100, args.rows // 50))
    estoque(out, rng)
    notas(out, rng)
    eventos(out)
    export_zip(out)
    catalogo(out)
    legado(out)
    reservadas(out)
    transacoes(out, rng, args.rows)
    trajetoria(out, rng)
    censo(out, rng)
    noise(out)
    print(f"inbox de demonstração gerada em {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
