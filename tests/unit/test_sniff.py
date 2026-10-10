import pandas as pd
import pytest
from openpyxl import Workbook

from datahack_ingest.catalog import Catalog
from datahack_ingest.extractors import ExtractState, build_extractor
from datahack_ingest.normalize import to_text_frame
from datahack_ingest.sniff import detect_delimiter, detect_encoding


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (b"a;b\n1;2\n", "utf-8"),
        (b"\xef\xbb\xbfa,b\n", "utf-8-sig"),
        ("nome\nJoão\n".encode("utf-16"), "utf-16"),
        ("nome;cidade\nJosé;Cuiabá\n".encode("cp1252"), "cp1252"),
        ("ação".encode()[:-1], "utf-8"),
    ],
)
def test_detect_encoding(raw, expected):
    assert detect_encoding(raw) == expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("id;valor\n1;10,5\n2;3,2\n", ";"),
        ("id,valor\n1,10\n", ","),
        ("id\tvalor\n1\t2\n", "\t"),
        ("id|valor\n1|2\n", "|"),
        ('id,obs\n1,"a;b;c"\n2,"d;e"\n', ","),
        ("unica\n1\n2\n", ","),
    ],
)
def test_detect_delimiter(text, expected):
    assert detect_delimiter(text) == expected


def _read(tmp_path, file_opts):
    spec = {"name": "f", "kind": "file", "chunk_size": 100, "file": file_opts}
    src = Catalog.model_validate({"sources": [spec]}).sources[0]
    ext = build_extractor(src, str(tmp_path))
    frames = [to_text_frame(f) for u in ext.units(ExtractState()) for f in u.frames]
    return pd.concat(frames, ignore_index=True)


def test_csv_auto_dialect_latin1_semicolon(tmp_path):
    raw = "Código;Descrição;Valor\n001;Pão de queijo;1.234,50\n002;Açaí;7,00\n"
    (tmp_path / "a.csv").write_bytes(raw.encode("cp1252"))
    df = _read(tmp_path, {"path": "a.csv", "format": "csv", "sep": "auto"})
    assert list(df.columns) == ["codigo", "descricao", "valor"]
    assert df["descricao"].tolist() == ["Pão de queijo", "Açaí"]
    assert df["codigo"].tolist() == ["001", "002"]


def test_csv_utf16_tab_export(tmp_path):
    (tmp_path / "a.txt").write_bytes("id\tnome\n1\tJoão\n".encode("utf-16"))
    df = _read(tmp_path, {"path": "a.txt", "format": "csv", "sep": "auto"})
    assert df.to_dict("records") == [{"id": "1", "nome": "João"}]


def test_directory_path_with_include_exclude_and_settle(tmp_path):
    folder = tmp_path / "pasta" / "sub"
    folder.mkdir(parents=True)
    (folder / "a.csv").write_text("id\n1\n", encoding="utf-8")
    (tmp_path / "pasta" / "b.CSV").write_text("id\n2\n", encoding="utf-8")
    (tmp_path / "pasta" / "~$lock.csv").write_text("id\n9\n", encoding="utf-8")
    (tmp_path / "pasta" / "leia.md").write_text("x", encoding="utf-8")
    df = _read(tmp_path, {"path": "pasta", "format": "csv", "include": ["*.csv"]})
    assert sorted(df["id"]) == ["1", "2"]
    spec = {
        "name": "f",
        "kind": "file",
        "file": {"path": "pasta", "format": "csv", "include": ["*.csv"], "min_age_seconds": 3600},
    }
    src = Catalog.model_validate({"sources": [spec]}).sources[0]
    ext = build_extractor(src, str(tmp_path))
    assert list(ext.units(ExtractState())) == []
    assert len(ext.deferred) == 2


def _inep_workbook(path, ano: int, notes: bool = True) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    for line in ("Ministério da Educação", "INEP", "", "Indicadores de Trajetória", "", "", "", ""):
        ws.append([line or None])
    ws.append(["CO_IES", "CO_CURSO", "NU_ANO_INGRESSO", "QT_INGRESSANTE", "TDA", None, None])
    ws.append(["0001", "0012345", ano, 120, 35.5])
    ws.append(["0002", "0099999", ano, 8, 12.0])
    if notes:
        ws.append([None])
        ws.append(["Fonte: Inep, Censo da Educação Superior."])
    wb.save(path)


def test_xlsx_header_on_row_9_and_footer_notes(tmp_path):
    _inep_workbook(tmp_path / "trajetoria_2015.xlsx", 2015)
    df = _read(
        tmp_path,
        {
            "path": "trajetoria_2015.xlsx",
            "format": "xlsx",
            "skip_rows": "auto",
            "drop_note_rows": True,
            "sheet_name": "auto",
        },
    )
    assert list(df.columns) == ["co_ies", "co_curso", "nu_ano_ingresso", "qt_ingressante", "tda"]
    assert df["co_curso"].tolist() == ["0012345", "0099999"]
    assert len(df) == 2


def test_csv_with_title_lines(tmp_path):
    raw = (
        "Relatório gerado em 01/10/2026\n\n"
        "CO_CURSO;QT_MAT;NO_CURSO\n001;10;Direito\n002;20;Medicina\n"
    )
    (tmp_path / "a.csv").write_bytes(raw.encode("cp1252"))
    df = _read(tmp_path, {"path": "a.csv", "format": "csv", "sep": "auto", "skip_rows": "auto"})
    assert list(df.columns) == ["co_curso", "qt_mat", "no_curso"]
    assert len(df) == 2
