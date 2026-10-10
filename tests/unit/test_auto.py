import dataclasses
import gzip
import io
import json
import sqlite3
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook

from datahack_ingest.catalog import Catalog
from datahack_ingest.discovery import classify, resolve_catalog
from datahack_ingest.extractors import ExtractState, build_extractor
from datahack_ingest.modelgen import (
    CHECKS,
    ColumnSpec,
    ColumnStats,
    TableSpec,
    _assign_output_names,
    classify_pii,
    cpf_is_valid,
    decide_type,
    key_candidates,
    render_gold,
    render_silver,
    render_yaml,
    singular,
)
from datahack_ingest.normalize import to_text_frame
from datahack_ingest.settings import Settings
from datahack_ingest.sniff import (
    detect_delimiter,
    detect_encoding,
    infer_json_records_path,
    infer_xml_record_tag,
)


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (b"a;b\n1;2\n", "utf-8"),
        (b"\xef\xbb\xbfa,b\n", "utf-8-sig"),
        ("nome\nJoão\n".encode("utf-16"), "utf-16"),
        ("nome;cidade\nJosé;Cuiabá\n".encode("cp1252"), "cp1252"),
        ("ação".encode()[:-1], "utf-8"),
    ],
)
def test_detect_encoding(raw, expected):
    assert detect_encoding(raw) == expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("id;valor\n1;10,5\n2;3,2\n", ";"),
        ("id,valor\n1,10\n", ","),
        ("id\tvalor\n1\t2\n", "\t"),
        ("id|valor\n1|2\n", "|"),
        ('id,obs\n1,"a;b;c"\n2,"d;e"\n', ","),
        ("unica\n1\n2\n", ","),
    ],
)
def test_detect_delimiter(text, expected):
    assert detect_delimiter(text) == expected


def test_infer_xml_record_tag_prefers_shallowest_repeated_record():
    xml = (
        b'<?xml version="1.0"?><lote xmlns="urn:x"><cab><v>1</v></cab>'
        b'<nota id="1"><item><q>1</q></item><item><q>2</q></item></nota>'
        b'<nota id="2"><item><q>3</q></item></nota></lote>'
    )
    assert infer_xml_record_tag(io.BytesIO(xml)) == "nota"


@pytest.mark.parametrize(
    ("payload", "expected"),
    [
        ([{"id": 1}], (True, None)),
        ({"meta": {"n": 1}, "data": {"items": [{"id": 1}]}}, (False, "data.items")),
        ({"tags": ["a"], "rows": [{"id": 1}]}, (False, "rows")),
        ({"id": 1}, (False, None)),
    ],
)
def test_infer_json_records_path(payload, expected):
    assert infer_json_records_path(io.BytesIO(json.dumps(payload).encode())) == expected


def _read(tmp_path, file_opts):
    spec = {"name": "f", "kind": "file", "chunk_size": 100, "file": file_opts}
    src = Catalog.model_validate({"sources": [spec]}).sources[0]
    ext = build_extractor(src, str(tmp_path))
    frames = [to_text_frame(f) for u in ext.units(ExtractState()) for f in u.frames]
    return pd.concat(frames, ignore_index=True)


def test_csv_auto_dialect_latin1_semicolon(tmp_path):
    raw = "Código;Descrição;Valor\n001;Pão de queijo;1.234,50\n002;Açaí;7,00\n"
    (tmp_path / "a.csv").write_bytes(raw.encode("cp1252"))
    df = _read(tmp_path, {"path": "a.csv", "format": "csv", "sep": "auto"})
    assert list(df.columns) == ["codigo", "descricao", "valor"]
    assert df["descricao"].tolist() == ["Pão de queijo", "Açaí"]
    assert df["codigo"].tolist() == ["001", "002"]


def test_csv_utf16_tab_export(tmp_path):
    (tmp_path / "a.txt").write_bytes("id\tnome\n1\tJoão\n".encode("utf-16"))
    df = _read(tmp_path, {"path": "a.txt", "format": "csv", "sep": "auto"})
    assert df.to_dict("records") == [{"id": "1", "nome": "João"}]


def test_directory_path_with_include_exclude_and_settle(tmp_path):
    folder = tmp_path / "pasta" / "sub"
    folder.mkdir(parents=True)
    (folder / "a.csv").write_text("id\n1\n", encoding="utf-8")
    (tmp_path / "pasta" / "b.CSV").write_text("id\n2\n", encoding="utf-8")
    (tmp_path / "pasta" / "~$lock.csv").write_text("id\n9\n", encoding="utf-8")
    (tmp_path / "pasta" / "leia.md").write_text("x", encoding="utf-8")
    df = _read(tmp_path, {"path": "pasta", "format": "csv", "include": ["*.csv"]})
    assert sorted(df["id"]) == ["1", "2"]
    spec = {
        "name": "f",
        "kind": "file",
        "file": {"path": "pasta", "format": "csv", "include": ["*.csv"], "min_age_seconds": 3600},
    }
    src = Catalog.model_validate({"sources": [spec]}).sources[0]
    ext = build_extractor(src, str(tmp_path))
    assert list(ext.units(ExtractState())) == []
    assert len(ext.deferred) == 2


@pytest.mark.parametrize(
    ("name", "fmt"),
    [
        ("Vendas 2026.CSV", "csv"),
        ("x.csv.gz", "csv"),
        ("x.ndjson.zst", "jsonl"),
        ("x.xlsx", "xlsx"),
        ("x.zip", "zip"),
        ("base.sqlite", "sqlite"),
        ("x.pdf", None),
        ("x.xls", None),
        ("semextensao", None),
    ],
)
def test_classify(name, fmt):
    assert classify(name)[0] == fmt


def _settings(tmp_path, **over):
    base = Settings.from_env()
    catalog = tmp_path / "sources.yml"
    catalog.write_text("version: 1\nsources: []\n", encoding="utf-8")
    return dataclasses.replace(
        base,
        catalog_path=catalog,
        landing_uri=str(tmp_path / "landing"),
        examples_enabled=False,
        inbox_settle_seconds=0,
        **over,
    )


def test_inbox_discovery_end_to_end(tmp_path):
    inbox = tmp_path / "landing" / "inbox"
    (inbox / "clientes").mkdir(parents=True)
    (inbox / "clientes" / "jan.csv").write_text("id;nome\n1;Ana\n", encoding="utf-8")
    (inbox / "clientes" / "fev.csv.gz").write_bytes(gzip.compress(b"id;nome\n2;Bia\n"))
    (inbox / "misto").mkdir()
    (inbox / "misto" / "a.csv").write_text("id\n1\n", encoding="utf-8")
    (inbox / "misto" / "b.json").write_text('[{"id": 1}]', encoding="utf-8")
    (inbox / "manual").mkdir()
    (inbox / "manual" / "dados.txt").write_text("001ANA\n", encoding="utf-8")
    (inbox / "manual" / "_source.yml").write_text(
        "name: posicional\nfile:\n  format: fixed_width\n  widths: [3, 3]\n  names: [cod, nome]\n",
        encoding="utf-8",
    )
    wb = Workbook()
    wb.active.title = "Produtos"
    wb.active.append(["id"])
    wb.create_sheet("Movimentos").append(["id"])
    wb.save(inbox / "Estoque Geral.xlsx")
    with zipfile.ZipFile(inbox / "lote.zip", "w") as zf:
        zf.writestr("x/a.xml", "<r><n><v>1</v></n></r>")
    with sqlite3.connect(inbox / "erp.db") as conn:
        conn.execute("CREATE TABLE pedidos (id integer, total real)")
        conn.execute("INSERT INTO pedidos VALUES (1, 9.5)")
    (inbox / "manual.pdf").write_bytes(b"%PDF")
    (inbox / "[ruim].csv").write_text("a\n1\n", encoding="utf-8")
    wb2 = Workbook()
    wb2.active.title = "Oculta"
    wb2.active.sheet_state = "hidden"
    wb2.create_sheet("Dados").append(["id"])
    wb2.save(inbox / "unica.xlsx")

    resolved = resolve_catalog(_settings(tmp_path))
    by_name = {s.name: s for s in resolved.catalog.sources}
    assert set(by_name) == {
        "clientes",
        "misto_csv",
        "misto_json",
        "posicional",
        "estoque_geral_produtos",
        "estoque_geral_movimentos",
        "a",
        "erp_pedidos",
        "unica",
    }
    assert by_name["unica"].file.sheet_name == "Dados"
    assert by_name["clientes"].file.include == ["*.csv", "*.csv.gz"]
    assert by_name["clientes"].file.sep == "auto"
    assert by_name["posicional"].file.format == "fixed_width"
    assert by_name["a"].file.format == "xml"
    assert by_name["a"].file.zip_members == ["x/a.xml"]
    assert by_name["erp_pedidos"].kind == "sql"
    assert all(resolved.origin(n) == "inbox" for n in by_name)
    reasons = {i["path"].rsplit("/", 1)[-1]: i["reason"] for i in resolved.ignored}
    assert "manual.pdf" in reasons and "[ruim].csv" in reasons

    rows = _read(tmp_path / "landing", by_name["clientes"].file.model_dump())
    assert sorted(rows["nome"]) == ["Ana", "Bia"]


def test_inbox_respects_catalog_names_and_bad_sidecar(tmp_path):
    inbox = tmp_path / "landing" / "inbox"
    (inbox / "vendas").mkdir(parents=True)
    (inbox / "vendas" / "a.csv").write_text("id\n1\n", encoding="utf-8")
    (inbox / "quebrada").mkdir()
    (inbox / "quebrada" / "a.csv").write_text("id\n1\n", encoding="utf-8")
    (inbox / "quebrada" / "_source.yml").write_text("load_strategy: merge\n", encoding="utf-8")
    settings = _settings(tmp_path)
    settings.catalog_path.write_text(
        "version: 1\nsources:\n  - name: vendas\n    kind: file\n"
        "    file: { path: x.csv, format: csv }\n",
        encoding="utf-8",
    )
    resolved = resolve_catalog(settings)
    names = {s.name for s in resolved.catalog.sources}
    assert names == {"vendas", "vendas_2"}
    assert resolved.errors and "primary_key" in resolved.errors[0]["reason"]


def test_inbox_can_be_disabled(tmp_path):
    (tmp_path / "landing" / "inbox").mkdir(parents=True)
    (tmp_path / "landing" / "inbox" / "a.csv").write_text("id\n1\n", encoding="utf-8")
    resolved = resolve_catalog(_settings(tmp_path, inbox_enabled=False))
    assert resolved.catalog.sources == []


def _stats(values: list[str | None]) -> ColumnStats:
    import re

    present = [v.strip() for v in values if v is not None and v.strip()]
    patterns = {
        "int_ok": r"^[+-]?\d{1,18}$",
        "lead0": r"^[+-]?0\d",
        "digits": r"^\d+$",
        "num_ok": r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$",
        "br_ok": r"^[+-]?(\d{1,3}(\.\d{3})+|\d+)(,\d+)?$",
        "br_comma": r",\d+$",
        "date_iso": r"^\d{4}-\d{2}-\d{2}$",
        "date_dmy": r"^(0?[1-9]|[12]\d|3[01])/(0?[1-9]|1[0-2])/\d{4}$",
        "date_mdy": r"^(0?[1-9]|1[0-2])/(0?[1-9]|[12]\d|3[01])/\d{4}$",
        "ts_iso": r"^\d{4}-\d{2}-\d{2}[ T]\d{1,2}:\d{2}",
        "ts_tz": r"^\d{4}-\d{2}-\d{2}[ T].*(Z|[+-]\d{2}(:?\d{2})?)$",
        "ts_dmy": r"^\d{1,2}/\d{1,2}/\d{4} \d{1,2}:\d{2}",
        "email": r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$",
        "cpf_fmt": r"^\d{3}\.\d{3}\.\d{3}-\d{2}$",
        "phone": r"^\(\d{2}\) ?9?\d{4}-\d{4}$",
    }
    counts = {k: sum(bool(re.search(p, v, re.I)) for v in present) for k, p in patterns.items()}
    counts["num_ok"] = sum(
        bool(re.search(patterns["num_ok"], v.replace("R$", "").replace(" ", ""))) for v in present
    )
    counts["br_ok"] = sum(
        bool(re.search(patterns["br_ok"], v.replace("R$", "").replace(" ", ""))) for v in present
    )
    counts["bool_ok"] = sum(v.lower() in {"sim", "não", "nao", "true", "false"} for v in present)
    counts["eleven"] = sum(len(re.sub(r"\D", "", v)) == 11 for v in present)
    counts["midnight"] = sum(bool(re.search(r"[ T]00:00(:00)?$", v)) for v in present)
    counts["json_ok"] = sum(v.startswith(("[", "{")) for v in present)
    assert set(counts) == set(CHECKS)
    return ColumnStats(
        rows=len(values),
        nn=len(present),
        nd=len(set(present)),
        maxlen=max((len(v) for v in present), default=0),
        counts=counts,
    )


@pytest.mark.parametrize(
    ("name", "values", "expected"),
    [
        ("quantidade", ["1", "2", "30"], ("bigint", None)),
        ("codigo", ["001", "002"], ("text", None)),
        ("chave_nfe", ["5" * 44, "4" * 44], ("text", None)),
        ("codigo_produto", ["10", "20"], ("text", None)),
        ("co_curso", ["12345", "678"], ("text", None)),
        ("tp_rede", ["1", "2"], ("text", None)),
        ("qt_ingressante", ["10", "200"], ("bigint", None)),
        ("valor", ["R$ 1.234,56", "7,00", "10"], ("numeric", ",")),
        ("preco", ["10.5", "3"], ("numeric", ".")),
        ("milhar", ["1.234", "2.000"], ("numeric", ".")),
        ("data", ["31/12/2025", "1/2/2026"], ("date", "dmy")),
        ("dia", ["2026-01-31"], ("date", "iso")),
        ("criado", ["2026-01-01T10:00:00Z"], ("timestamptz", "iso")),
        ("hora", ["2026-01-01 10:00:00"], ("timestamp", "iso")),
        ("ativo", ["Sim", "Não", "sim"], ("boolean", None)),
        ("misto", ["1", "2", "x"], ("text", None)),
        ("dia_excel", ["2026-01-01 00:00:00", "2026-02-01 00:00:00"], ("date", "iso")),
        ("itens", ['[{"a": 1}]', '{"b": 2}'], ("jsonb", None)),
        ("vazio", [None, " "], ("text", None)),
    ],
)
def test_decide_type(name, values, expected):
    kind, fmt, _ = decide_type(name, _stats(values), 0.98)
    assert (kind, fmt) == expected


def test_cpf_checksum():
    assert cpf_is_valid("529.982.247-25")
    assert cpf_is_valid("52998224725")
    assert not cpf_is_valid("529.982.247-24")
    assert not cpf_is_valid("111.111.111-11")


@pytest.mark.parametrize(
    ("name", "values", "cpf_ratio", "declared", "expected"),
    [
        ("cpf_cliente", ["x"], None, False, ("identificador", True)),
        ("contato", ["a@b.com", "c@d.org"], None, False, ("identificador", False)),
        ("documento", ["529.982.247-25"], None, False, ("identificador", True)),
        ("doc", ["52998224725"], 1.0, False, ("identificador", True)),
        ("fone", ["(65) 99999-0000"], None, False, ("identificador", True)),
        ("nome_cliente", ["Ana"], None, False, ("pessoal", False)),
        ("nome_produto", ["Arroz"], None, False, (None, False)),
        ("segredo", ["x"], None, True, ("identificador", False)),
    ],
)
def test_classify_pii(name, values, cpf_ratio, declared, expected):
    assert classify_pii(name, _stats(values), cpf_ratio, declared) == expected


def _spec(materialization="table", dedup="key"):
    cols = [
        ColumnSpec(name="id", ordinal=1, inferred_type="bigint", is_key=dedup == "key"),
        ColumnSpec(name="cpf", ordinal=2, pii_class="identificador", hash_digits=True),
        ColumnSpec(name="valor", ordinal=3, inferred_type="numeric", type_format=","),
        ColumnSpec(name="order", ordinal=4),
    ]
    _assign_output_names(cols)
    spec = TableSpec(
        table="clientes",
        source="clientes",
        columns=cols,
        key_columns=["id"] if dedup == "key" else [],
        dedup=dedup,
        materialization=materialization,
    )
    spec.silver_alias = spec.gold_alias = "clientes"
    return spec


def test_render_silver_table_and_incremental():
    sql_table = render_silver(_spec())
    assert "materialized='table'" in sql_table
    assert "partition by nullif(btrim(\"id\"), '')" in sql_table
    assert 'dh_hash_pii(\'"cpf"\', digits_only=true) }} as "cpf_hash"' in sql_table
    assert "dh_to_numeric('\"valor\"', decimal=',')" in sql_table
    assert 'dh_clean_text(\'"order"\') }} as "order"' in sql_table
    assert "is_incremental" not in sql_table

    sql_inc = render_silver(_spec("incremental", "row_hash"))
    assert "unique_key=['_dh_row_hash']" in sql_inc
    assert "partition by _dh_row_hash" in sql_inc
    assert "{% if is_incremental() %}\n    where _dh_ingested_at" in sql_inc


def test_render_gold_and_yaml():
    gold = render_gold(_spec())
    assert "ref('auto_silver__clientes')" in gold
    assert "where _dh_deleted_at is null" in gold
    assert '"cpf_hash"' in gold and '"cpf"' not in gold.replace('"cpf_hash"', "")
    sources, models = render_yaml([_spec()])
    assert "bronze_auto" in sources and "severity: warn" in sources
    assert "auto_silver__clientes" in models and "auto_gold__clientes" in models
    assert "dh_invalid_ratio" in models


@pytest.mark.parametrize(
    ("word", "expected"),
    [("transacoes", "transacao"), ("clientes", "cliente"), ("fiscais", "fiscal"), ("id", "id")],
)
def test_singular(word, expected):
    assert singular(word) == expected


@pytest.mark.parametrize(
    ("table", "columns", "expected"),
    [
        ("transacoes", ["valor", "transacao_id"], ["transacao_id"]),
        ("pedidos", ["cliente_id", "id", "pedido_id"], ["id", "pedido_id"]),
        ("itens_venda", ["venda_id", "produto_id"], ["venda_id"]),
        ("reservadas", ["order", "select"], []),
        ("clientes", ["codigo", "nome"], ["codigo"]),
        ("notas_fiscais", ["nota_id"], ["nota_id"]),
    ],
)
def test_key_candidates_only_own_identifiers(table, columns, expected):
    assert key_candidates(table, columns) == expected


def _inep_workbook(path, ano: int, notes: bool = True) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    for line in ("Ministério da Educação", "INEP", "", "Indicadores de Trajetória", "", "", "", ""):
        ws.append([line or None])
    ws.append(["CO_IES", "CO_CURSO", "NU_ANO_INGRESSO", "QT_INGRESSANTE", "TDA", None, None])
    ws.append(["0001", "0012345", ano, 120, 35.5])
    ws.append(["0002", "0099999", ano, 8, 12.0])
    if notes:
        ws.append([None])
        ws.append(["Fonte: Inep, Censo da Educação Superior."])
    wb.save(path)


def test_xlsx_header_on_row_9_and_footer_notes(tmp_path):
    _inep_workbook(tmp_path / "trajetoria_2015.xlsx", 2015)
    df = _read(
        tmp_path,
        {
            "path": "trajetoria_2015.xlsx",
            "format": "xlsx",
            "skip_rows": "auto",
            "drop_note_rows": True,
            "sheet_name": "auto",
        },
    )
    assert list(df.columns) == ["co_ies", "co_curso", "nu_ano_ingresso", "qt_ingressante", "tda"]
    assert df["co_curso"].tolist() == ["0012345", "0099999"]
    assert len(df) == 2


def test_csv_with_title_lines(tmp_path):
    raw = (
        "Relatório gerado em 01/10/2026\n\n"
        "CO_CURSO;QT_MAT;NO_CURSO\n001;10;Direito\n002;20;Medicina\n"
    )
    (tmp_path / "a.csv").write_bytes(raw.encode("cp1252"))
    df = _read(tmp_path, {"path": "a.csv", "format": "csv", "sep": "auto", "skip_rows": "auto"})
    assert list(df.columns) == ["co_curso", "qt_mat", "no_curso"]
    assert len(df) == 2


def test_family_groups_years_and_splits_zip_tables(tmp_path):
    inbox = tmp_path / "landing" / "inbox"
    inbox.mkdir(parents=True)
    for ano in (2021, 2022):
        _inep_workbook(inbox / f"CPC_{ano}.xlsx", ano)
    for ano in (2023, 2024):
        with zipfile.ZipFile(inbox / f"microdados_censo_{ano}.zip", "w") as zf:
            zf.writestr(
                f"m{ano}/dados/MICRODADOS_CADASTRO_CURSOS_{ano}.CSV", "CO_CURSO;QT_MAT\n1;2\n"
            )
            zf.writestr(f"m{ano}/dados/MICRODADOS_ED_SUP_IES_{ano}.CSV", "CO_IES;NO_IES\n1;X\n")
            zf.writestr(f"m{ano}/Anexos/dicionário_dados_{ano}.xlsx", b"x")
            zf.writestr(f"m{ano}/leia-me/leia-me.pdf", b"%PDF")
    (inbox / "dicionario_trajetoria.xlsx").write_bytes(b"x")
    (inbox / "_sources.yml").write_text(
        "microdados_cadastro_cursos:\n"
        "  transforms:\n"
        "    - { op: filter, column: qt_mat, operator: eq, value: '2' }\n",
        encoding="utf-8",
    )

    resolved = resolve_catalog(_settings(tmp_path))
    by_name = {s.name: s for s in resolved.catalog.sources}
    assert set(by_name) == {"cpc", "microdados_cadastro_cursos", "microdados_ed_sup_ies"}
    assert by_name["cpc"].file.include == ["CPC_2021.xlsx", "CPC_2022.xlsx"]
    assert by_name["cpc"].file.sheet_name == "auto"
    cursos = by_name["microdados_cadastro_cursos"].file
    assert cursos.include == ["microdados_censo_2023.zip", "microdados_censo_2024.zip"]
    assert len(cursos.zip_members) == 2 and cursos.recursive is False
    assert by_name["microdados_cadastro_cursos"].transforms[0].column == "qt_mat"
    reasons = " ".join(i["reason"] for i in resolved.ignored)
    assert "documentação" in reasons

    landing = tmp_path / "landing"
    assert sorted(_read(landing, cursos.model_dump())["co_curso"]) == ["1", "1"]
    cpc = _read(landing, by_name["cpc"].file.model_dump())
    assert sorted(cpc["nu_ano_ingresso"]) == ["2021", "2021", "2022", "2022"]
