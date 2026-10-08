import bz2
import gzip
import json
import zipfile

import fastavro
import pandas as pd
import pyarrow as pa
import pytest
import zstandard
from openpyxl import Workbook
from pyarrow import orc

from datahack_ingest.catalog import Catalog
from datahack_ingest.extractors import ExtractState, build_extractor
from datahack_ingest.normalize import to_text_frame

CSV = "id,nome\n1,Ana\n2,Bia\n3,\n"


def _read(tmp_path, file_opts: dict, chunk_size: int = 100) -> pd.DataFrame:
    spec = {"name": "f", "kind": "file", "chunk_size": chunk_size, "file": file_opts}
    src = Catalog.model_validate({"sources": [spec]}).sources[0]
    ext = build_extractor(src, str(tmp_path))
    frames = [to_text_frame(f) for u in ext.units(ExtractState()) for f in u.frames]
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


@pytest.mark.parametrize(
    ("name", "writer"),
    [
        ("a.csv.gz", lambda p: p.write_bytes(gzip.compress(CSV.encode()))),
        ("a.csv.bz2", lambda p: p.write_bytes(bz2.compress(CSV.encode()))),
        ("a.csv.zst", lambda p: p.write_bytes(zstandard.ZstdCompressor().compress(CSV.encode()))),
    ],
)
def test_compressed_csv(tmp_path, name, writer):
    writer(tmp_path / name)
    df = _read(tmp_path, {"path": name, "format": "csv"})
    assert df["nome"].tolist() == ["Ana", "Bia", None]


def test_zip_with_member_pattern(tmp_path):
    with zipfile.ZipFile(tmp_path / "lote.zip", "w") as zf:
        zf.writestr("2026/jan.csv", "id\n1\n2\n")
        zf.writestr("2026/fev.csv", "id\n3\n")
        zf.writestr("leia-me.txt", "ignorar")
    df = _read(tmp_path, {"path": "lote.zip", "format": "csv", "zip_member_pattern": "*.csv"})
    assert sorted(df["id"].tolist()) == ["1", "2", "3"]


def test_xml_streaming_with_attributes_and_namespace(tmp_path):
    (tmp_path / "nf.xml").write_text(
        '<?xml version="1.0"?><lote xmlns="http://nfe"><nota numero="10"><valor>99.5</valor>'
        "<emitente><cnpj>123</cnpj></emitente><item>A</item><item>B</item></nota>"
        '<nota numero="11"><valor>1</valor></nota></lote>',
        encoding="utf-8",
    )
    df = _read(tmp_path, {"path": "nf.xml", "format": "xml", "record_tag": "nota"}, chunk_size=100)
    assert df["numero"].tolist() == ["10", "11"]
    assert df["valor"].tolist() == ["99.5", "1"]
    assert df["emitente_cnpj"].tolist()[0] == "123"
    assert json.loads(df["item"].tolist()[0]) == ["A", "B"]


def test_xml_rejects_entity_expansion(tmp_path):
    (tmp_path / "bomb.xml").write_text(
        '<?xml version="1.0"?><!DOCTYPE r [<!ENTITY a "aaaa">]><r><nota>&a;</nota></r>',
        encoding="utf-8",
    )
    with pytest.raises(Exception, match=r"(?i)entit"):
        _read(tmp_path, {"path": "bomb.xml", "format": "xml", "record_tag": "nota"})


def test_fixed_width(tmp_path):
    (tmp_path / "a.txt").write_text("001ANA       0100\n002BIA       0250\n", encoding="utf-8")
    df = _read(
        tmp_path,
        {
            "path": "a.txt",
            "format": "fixed_width",
            "widths": [3, 10, 4],
            "names": ["id", "nome", "valor"],
        },
    )
    assert df.to_dict("records")[1] == {"id": "002", "nome": "BIA", "valor": "0250"}


def test_avro(tmp_path):
    schema = {
        "type": "record",
        "name": "Venda",
        "fields": [
            {"name": "id", "type": "long"},
            {
                "name": "cliente",
                "type": {
                    "type": "record",
                    "name": "C",
                    "fields": [{"name": "uf", "type": "string"}],
                },
            },
            {"name": "obs", "type": ["null", "string"], "default": None},
        ],
    }
    with open(tmp_path / "v.avro", "wb") as f:
        fastavro.writer(
            f,
            schema,
            [{"id": 1, "cliente": {"uf": "MT"}}, {"id": 2, "cliente": {"uf": "SP"}, "obs": "x"}],
        )
    df = _read(tmp_path, {"path": "v.avro", "format": "avro"})
    assert df["id"].tolist() == ["1", "2"]
    assert df["cliente_uf"].tolist() == ["MT", "SP"]
    assert df["obs"].tolist() == [None, "x"]


def test_orc(tmp_path):
    orc.write_table(pa.table({"id": pa.array([1, None, 3], pa.int64())}), str(tmp_path / "a.orc"))
    df = _read(tmp_path, {"path": "a.orc", "format": "orc"})
    assert df["id"].tolist() == ["1", None, "3"]


def test_xlsx_streaming_in_chunks(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Código", "Valor"])
    for i in range(250):
        ws.append([i, i * 1.5])
    ws.append([None, None])
    wb.save(tmp_path / "a.xlsx")
    df = _read(tmp_path, {"path": "a.xlsx", "format": "xlsx"}, chunk_size=100)
    assert len(df) == 250
    assert list(df.columns) == ["codigo", "valor"]
    assert df["codigo"].tolist()[:2] == ["0", "1"]
    assert df["valor"].tolist()[1] == "1.5"


@pytest.mark.parametrize(
    ("payload", "opts", "expected"),
    [
        ([{"id": 1}, {"id": 2.5}], {}, ["1", "2.5"]),
        ({"data": {"items": [{"id": 7}]}}, {"records_path": "data.items"}, ["7"]),
        ({"id": 9}, {}, ["9"]),
    ],
)
def test_json_streaming_variants(tmp_path, payload, opts, expected):
    (tmp_path / "a.json").write_text(json.dumps(payload), encoding="utf-8")
    df = _read(tmp_path, {"path": "a.json", "format": "json", **opts})
    assert df["id"].tolist() == expected


def test_format_specific_validation():
    with pytest.raises(Exception, match="record_tag"):
        Catalog.model_validate(
            {"sources": [{"name": "x", "kind": "file", "file": {"path": "a", "format": "xml"}}]}
        )
    with pytest.raises(Exception, match="widths"):
        Catalog.model_validate(
            {
                "sources": [
                    {"name": "x", "kind": "file", "file": {"path": "a", "format": "fixed_width"}}
                ]
            }
        )
