from __future__ import annotations

import hashlib
import io
import logging
import os
import posixpath
import time
import zipfile
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from datetime import datetime
from fnmatch import fnmatch
from itertools import chain, islice
from typing import IO, Any

import fsspec
import pandas as pd
from fsspec.utils import infer_compression
from openpyxl import load_workbook

from datahack_ingest.catalog import FileSource
from datahack_ingest.extractors.base import ExtractState, ExtractUnit
from datahack_ingest.normalize import to_text
from datahack_ingest.sniff import (
    SAMPLE_BYTES,
    decode_sample,
    detect_csv_header_row,
    detect_delimiter,
    detect_encoding,
    detect_header_row,
)

log = logging.getLogger(__name__)
_HASH_BLOCK = 8 * 1024 * 1024
Opener = Callable[[], Any]


def resolve_uri(landing_uri: str, path: str) -> str:
    if "://" in path or os.path.isabs(path):
        return path
    return f"{landing_uri.rstrip('/')}/{path.lstrip('/')}"


def _seekable(f: IO[bytes]) -> IO[bytes]:
    try:
        if f.seekable():
            return f
    except (AttributeError, OSError):
        pass
    return io.BytesIO(f.read())


def _modified_epoch(info: dict[str, Any]) -> float | None:
    for key in ("mtime", "LastModified", "last_modified", "updated", "modified"):
        value = info.get(key)
        if isinstance(value, int | float):
            return float(value)
        if isinstance(value, datetime):
            return value.timestamp()
    return None


def accepts(name: str, include: list[str], exclude: list[str]) -> bool:
    base = posixpath.basename(name.rstrip("/")).lower()
    if any(fnmatch(base, pattern.lower()) for pattern in exclude):
        return False
    return any(fnmatch(base, pattern.lower()) for pattern in include)


def _chunks(iterable, size: int) -> Iterator[list]:
    it = iter(iterable)
    while batch := list(islice(it, size)):
        yield batch


class FileExtractor:
    def __init__(self, source: FileSource, landing_uri: str) -> None:
        self.source = source
        self.opts = source.file
        self.urlpath = resolve_uri(landing_uri, self.opts.path)
        self.skipped = 0
        self.deferred: list[str] = []

    def list_files(self) -> tuple[fsspec.AbstractFileSystem, list[str]]:
        fs, _, paths = fsspec.get_fs_token_paths(self.urlpath)
        candidates: set[str] = set()
        for p in paths:
            if fs.isdir(p) and self.opts.recursive:
                candidates.update(f for f in fs.find(p) if not f.endswith("/"))
            elif fs.isdir(p):
                candidates.update(f for f in fs.ls(p, detail=False) if fs.isfile(f))
            elif fs.isfile(p):
                candidates.add(p)
        files = sorted(p for p in candidates if accepts(p, self.opts.include, self.opts.exclude))
        if self.opts.min_age_seconds:
            files = self._settled(fs, files)
        return fs, files

    def _settled(self, fs: fsspec.AbstractFileSystem, files: list[str]) -> list[str]:
        now = time.time()
        ready = []
        for path in files:
            modified = _modified_epoch(fs.info(path))
            if modified is not None and now - modified < self.opts.min_age_seconds:
                self.deferred.append(path)
                log.info("arquivo recente aguardando estabilizar", extra={"file": path})
                continue
            ready.append(path)
        return ready

    def units(self, state: ExtractState) -> Iterator[ExtractUnit]:
        fs, files = self.list_files()
        if not files:
            if state.strategy == "full":
                raise FileNotFoundError(
                    f"{self.source.name}: nenhum arquivo em {self.urlpath}; "
                    "carga full abortada para não esvaziar a Bronze."
                )
            log.warning("nenhum arquivo encontrado", extra={"urlpath": self.urlpath})
            return
        for path in files:
            digest, size = self._digest(fs, path)
            if state.strategy != "full" and digest in state.loaded_file_hashes:
                log.info("arquivo já carregado (idempotência)", extra={"file": path})
                self.skipped += 1
                continue
            yield ExtractUnit(
                key=fs.unstrip_protocol(path),
                frames=self._read(fs, path),
                file_sha256=digest,
                file_size=size,
            )

    @staticmethod
    def _digest(fs: fsspec.AbstractFileSystem, path: str) -> tuple[str, int]:
        h = hashlib.sha256()
        size = 0
        with fs.open(path, "rb") as f:
            while block := f.read(_HASH_BLOCK):
                h.update(block)
                size += len(block)
        return h.hexdigest(), size

    def _is_zip(self, path: str) -> bool:
        comp = self.opts.compression
        return comp == "zip" or (comp == "infer" and path.lower().endswith(".zip"))

    def _compression(self, path: str) -> str | None:
        comp = self.opts.compression
        if comp == "none":
            return None
        if comp == "infer":
            return infer_compression(path)
        return comp

    def _openers(self, fs: fsspec.AbstractFileSystem, path: str) -> Iterator[tuple[str, Opener]]:
        if not self._is_zip(path):
            compression = self._compression(path)
            yield path, lambda: fs.open(path, "rb", compression=compression)
            return
        wanted = {m.lower() for m in self.opts.zip_members}
        with fs.open(path, "rb") as raw:
            members = sorted(
                n
                for n in zipfile.ZipFile(raw).namelist()
                if not n.endswith("/")
                and fnmatch(n.lower(), self.opts.zip_member_pattern.lower())
                and accepts(n, ["*"], self.opts.exclude)
                and (not wanted or n.lower() in wanted)
            )
        if not members:
            raise FileNotFoundError(
                f"{self.source.name}: zip {path} sem membros para '{self.opts.zip_member_pattern}'"
            )
        for member in members:

            @contextmanager
            def opener(m=member):
                with (
                    fs.open(path, "rb") as raw_zip,
                    zipfile.ZipFile(raw_zip) as zf,
                    zf.open(m) as f,
                ):
                    yield f

            yield f"{path}!{member}", opener

    def _read(self, fs: fsspec.AbstractFileSystem, path: str) -> Iterator[pd.DataFrame]:
        reader = {"csv": self._read_csv, "xlsx": self._read_xlsx}[self.opts.format]
        for _, opener in self._openers(fs, path):
            yield from reader(opener)

    def _sample(self, opener: Opener) -> bytes:
        with opener() as f:
            return f.read(SAMPLE_BYTES)

    def _encoding(self, opener: Opener, sample: bytes | None = None) -> str:
        if self.opts.encoding != "auto":
            return self.opts.encoding
        return detect_encoding(sample if sample is not None else self._sample(opener))

    def _csv_dialect(self, opener: Opener) -> tuple[str, str, int]:
        sample = None
        auto_skip = self.opts.skip_rows == "auto"
        if self.opts.encoding == "auto" or self.opts.sep == "auto" or auto_skip:
            sample = self._sample(opener)
        encoding = self._encoding(opener, sample)
        text = decode_sample(sample, encoding) if sample is not None else ""
        sep = self.opts.sep
        if sep == "auto":
            sep = detect_delimiter(text, self.opts.quotechar)
        elif sep == r"\t":
            sep = "\t"
        skip = detect_csv_header_row(text, sep, self.opts.quotechar) if auto_skip else 0
        return sep, encoding, skip if auto_skip else int(self.opts.skip_rows)

    def _read_csv(self, opener: Opener) -> Iterator[pd.DataFrame]:
        sep, encoding, skip = self._csv_dialect(opener)
        log.info("dialeto csv", extra={"sep": sep, "encoding": encoding, "skip_rows": skip})
        with opener() as f:
            yield from pd.read_csv(
                f,
                sep=sep,
                encoding=encoding,
                encoding_errors="replace",
                quotechar=self.opts.quotechar,
                skiprows=skip,
                dtype=str,
                keep_default_na=False,
                na_values=[""],
                chunksize=self.source.chunk_size,
            )

    def _worksheet(self, wb):
        sheet = self.opts.sheet_name
        if sheet == "auto":
            visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
            return visible[0] if visible else wb.worksheets[0]
        return wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]

    def _xlsx_rows(self, ws) -> tuple[tuple | None, Iterator[tuple]]:
        rows = ws.iter_rows(values_only=True)
        if self.opts.skip_rows == "auto":
            head = list(islice(rows, 60))
            if not head:
                return None, iter(())
            idx = detect_header_row(head)
            if idx:
                log.info("cabeçalho detectado", extra={"header_row": idx + 1})
            return head[idx], chain(head[idx + 1 :], rows)
        for _ in range(int(self.opts.skip_rows)):
            next(rows, None)
        return next(rows, None), rows

    def _read_xlsx(self, opener: Opener) -> Iterator[pd.DataFrame]:
        with opener() as f:
            wb = load_workbook(_seekable(f), read_only=True, data_only=True)
            try:
                header, rows = self._xlsx_rows(self._worksheet(wb))
                if header is None:
                    return
                filled = [i for i, h in enumerate(header) if h is not None and str(h).strip()]
                width = (filled[-1] + 1) if filled else len(header)
                columns = [
                    str(h) if h is not None and str(h).strip() else f"col_{i + 1}"
                    for i, h in enumerate(header[:width])
                ]
                min_cells = 2 if self.opts.drop_note_rows and width >= 4 else 1
                dropped = 0
                for batch in _chunks(rows, self.source.chunk_size):
                    data = []
                    for r in batch:
                        values = (list(r) + [None] * width)[:width]
                        cells = sum(1 for v in values if v is not None and str(v).strip())
                        if cells < min_cells:
                            dropped += cells > 0
                            continue
                        data.append([to_text(v) for v in values])
                    if data:
                        yield pd.DataFrame(data, columns=columns, dtype=object)
                if dropped:
                    log.info("linhas de nota descartadas", extra={"rows": dropped})
            finally:
                wb.close()
